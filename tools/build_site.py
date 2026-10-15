"""Build the generated parts of the DARIA website.

    python tools/build_site.py            # everything
    python tools/build_site.py deck       # portfolio slides + PDF, from the PowerPoint
    python tools/build_site.py images     # stills and product cut-outs on the home page
    python tools/build_site.py video      # web films
    python tools/build_site.py frames     # stills taken from the films
    python tools/build_site.py og         # the social share image
    python tools/build_site.py icons      # favicon, app icons and web manifest

The source of truth is `DARIA AI Presentation 2026.pptx`. The `deck` step
opens it in PowerPoint (read-only), exports every slide and the PDF, and
rewrites the slide list inside site/portfolio.html, so the portfolio page and
the download always match the deck. Run it again after editing the deck.

Needs:  pip install Pillow imageio-ffmpeg PyMuPDF python-pptx
        PowerPoint (for `deck`) and Chrome or Edge (for `og`).
"""

import html
import io
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
PPTX = ROOT / "DARIA AI Presentation 2026.pptx"
SRC = ROOT / "Projects"
SITE = ROOT / "site"
IMG = SITE / "assets" / "img"
VID = SITE / "assets" / "video"
DECK = SITE / "assets" / "deck"
PDF = SITE / "downloads" / "DARIA-AI-Creative-Studio-Portfolio.pdf"
PORTFOLIO = SITE / "portfolio.html"

# ----------------------------------------------------------------- films ---

# slug: (name, master). The deck embeds these same files; they are matched
# to the films on its slides by file size.
FILMS = {
    "showreel": ("Showreel", "Showreel.mp4"),
    "clarins": ("Clarins", "Clarins/Final/Product Ad.mp4"),
    "dreamies": ("Dreamies", "Dreamies/Final/Dreamies - Product Ad.mp4"),
    "havaianas": ("Havaianas", "Havaianas/Final/Havaianas - Brand Ad.mp4"),
    "escudo": ("Escudo", "Escudo/Final/Escudo - Brand Ad.mp4"),
    "blue-banana": ("Blue Banana", "Blue Banana/Final/Blue Banana - Brand Ad.mp4"),
    "manucurist": ("Manucurist", "Manucurist/Final/Manucurist - Product Ad.mp4"),
}

# The home page opens on this film, so it also gets a lighter cut for phones.
HERO_FILM = "showreel"

# One still per film, taken at this many seconds in. The home page uses them
# as the frame a film shows before it plays, and two of them in the stills row.
FRAMES = {
    "showreel": 0.5,
    "clarins": 0.6,
    "dreamies": 6.4,
    "havaianas": 3.8,
    "escudo": 14.8,
    "blue-banana": 3.6,
    "manucurist": 5.3,
}

# ---------------------------------------------------------------- stills ---

# Finals the home page shows. Each gets a -sm and a -lg WebP.
FINALS = {
    "desigual": [
        ("look1-1", "Desigual/Final/Flower Dress, Sandals and Backpack (1).png"),
        ("look1-2", "Desigual/Final/Flower Dress, Sandals and Backpack (2).png"),
        ("look1-3", "Desigual/Final/Flower Dress, Sandals and Backpack (3).png"),
        ("look3-2", "Desigual/Final/Turqoise Dress and Boots (2).png"),
    ],
    "chilli-beans": [(f"final-{n}", f"Chilli Beans/Final/Product Ad ({n}).png") for n in range(1, 4)],
    "toothpaste": [("final-1", "Toothpaste/Final/Product Ad (1).png")],
    "cookies": [("final-1", "Cookies/Final/Cookie Ad (1).png")],
}

# Product cut-outs come out of the deck itself, with the crops and the
# transparent backgrounds made there. Keyed by the title on the case slide.
REF_SLIDES = {
    "Desigual — Look 01": "desigual",
    "Chilli Beans": "chilli-beans",
    "Escudo": "escudo",
    "Blue Banana": "blue-banana",
    "Clarins": "clarins",
    "Havaianas": "havaianas",
    "Dreamies": "dreamies",
    "Manucurist": "manucurist",
}
REF_TRAY_FROM = 0.72  # cut-outs sit in the tray on the right of the slide

PRIVATE = re.compile(r"\S+@\S+\.\w+|\+?\d[\d ]{7,}\d")  # never publish as text


def flat(im):
    """Drop alpha onto white."""
    if im.mode in ("RGBA", "LA", "P"):
        im = im.convert("RGBA")
        bg = Image.new("RGB", im.size, "#FFFFFF")
        bg.paste(im, mask=im.split()[-1])
        return bg
    return im.convert("RGB")


def fit_width(im, width):
    if im.width <= width:
        return im
    return im.resize((width, round(im.height * width / im.width)), Image.LANCZOS)


def clear(folder, pattern="*"):
    folder.mkdir(parents=True, exist_ok=True)
    for f in folder.glob(pattern):
        if f.is_file():
            f.unlink()


def deck():
    from pptx import Presentation
    return Presentation(PPTX)


def tray_pictures(pres, slide):
    return [s for s in slide.shapes if s.shape_type == 13 and not s.element.xpath(".//a:videoFile")
            and s.left / pres.slide_width > REF_TRAY_FROM]


def build_refs():
    out = IMG / "refs"
    clear(out)
    pres = deck()
    found = set()
    for slide in pres.slides:
        pics = tray_pictures(pres, slide)
        lines = {s.text_frame.text.strip() for s in slide.shapes if s.has_text_frame}
        slug = next((v for k, v in REF_SLIDES.items() if k in lines), None)
        if not slug or not pics:
            continue
        found.add(slug)
        for n, shape in enumerate(sorted(pics, key=lambda s: (s.top, s.left)), 1):
            im = Image.open(io.BytesIO(shape.image.blob)).convert("RGBA")
            w, h = im.size
            im = im.crop((round(w * shape.crop_left), round(h * shape.crop_top),
                          round(w * (1 - shape.crop_right)), round(h * (1 - shape.crop_bottom))))
            im = im.crop(im.getbbox())  # trim the empty margin so cut-outs sit evenly
            im.thumbnail((520, 520), Image.LANCZOS)
            im.save(out / f"{slug}-{n}.webp", "WEBP", quality=88, method=6)
        print("refs  ", slug, len(pics))
    missing = set(REF_SLIDES.values()) - found
    assert not missing, f"no case slide found for: {', '.join(sorted(missing))}"


def build_images():
    if IMG.exists():
        shutil.rmtree(IMG)  # everything in here is generated
    for case, items in FINALS.items():
        out = IMG / case
        clear(out)
        for stem, master in items:
            im = flat(Image.open(SRC / master))
            fit_width(im, 640).save(out / f"{stem}-sm.webp", "WEBP", quality=80, method=6)
            fit_width(im, 1280).save(out / f"{stem}-lg.webp", "WEBP", quality=82, method=6)
            print("still ", case, stem)
    build_refs()


def ffmpeg():
    import imageio_ffmpeg
    return imageio_ffmpeg.get_ffmpeg_exe()


def encode(master, out, *extra):
    subprocess.run([ffmpeg(), "-v", "error", "-y", "-i", str(SRC / master), *extra,
                    "-c:v", "libx264", "-preset", "slow", "-pix_fmt", "yuv420p",
                    "-c:a", "aac", "-movflags", "+faststart", str(out)], check=True)
    print("film  ", out.name, f"{out.stat().st_size / 1e6:.1f} MB")


def build_frames():
    clear(VID, "*.webp")
    for slug, at in FRAMES.items():
        with tempfile.TemporaryDirectory() as tmp:
            png = Path(tmp) / "frame.png"
            subprocess.run([ffmpeg(), "-v", "error", "-y", "-ss", str(at), "-i", str(SRC / FILMS[slug][1]),
                            "-frames:v", "1", str(png)], check=True)
            frame = Image.open(png).convert("RGB")
        fit_width(frame, 900 if frame.height > frame.width else 1600).save(
            VID / f"{slug}-frame.webp", "WEBP", quality=82, method=6)
        print("frame ", slug)


def build_video():
    clear(VID, "*.mp4")
    # Masters run at ~25 Mbps. CRF 24 capped at 4 Mbps is visually clean at
    # 1080p and brings 390 MB down to a size a page can carry.
    for slug, (_, master) in FILMS.items():
        encode(master, VID / f"{slug}.mp4", "-crf", "24", "-maxrate", "4M", "-bufsize", "8M", "-b:a", "128k")
    encode(FILMS[HERO_FILM][1], VID / f"{HERO_FILM}-720.mp4", "-vf", "scale=-2:720",
           "-crf", "26", "-maxrate", "1800k", "-bufsize", "3600k", "-b:a", "96k")
    build_frames()


# ------------------------------------------------------------------ deck ---

# PowerPoint does the export, so slides and PDF look as they do in the deck.
# Opened read-only and windowless; an already-open PowerPoint is left running.
PS_EXPORT = r"""
$ErrorActionPreference = 'Stop'
$wasRunning = [bool](Get-Process POWERPNT -ErrorAction SilentlyContinue)
$ppt = New-Object -ComObject PowerPoint.Application
try {
  $pres = $ppt.Presentations.Open($env:DARIA_PPTX, -1, 0, 0)
  $pres.Export((Join-Path $env:DARIA_OUT 'png'), 'PNG', 2560, 1440)
  $pres.SaveAs((Join-Path $env:DARIA_OUT 'deck.pdf'), 32)
  $pres.Close()
} finally {
  if (-not $wasRunning) { $ppt.Quit() }
}
"""

PLAY = '<span class="play"><svg viewBox="0 0 24 24" aria-hidden="true"><path d="M8 5v14l11-7z"/></svg></span>'
FOOTER_NOISE = re.compile(r"^(\d\d / \d\d|DARIA — AI CREATIVE STUDIO|FINAL — AI GENERATED|→|\d\d)$")


def read_deck():
    """Per slide: its text (for alt text) and where each film sits."""
    by_size = {(SRC / master).stat().st_size: slug for slug, (_, master) in FILMS.items()}
    with zipfile.ZipFile(PPTX) as z:
        part_size = {"/" + i.filename: i.file_size for i in z.infolist()}

    pres = deck()
    w, h = pres.slide_width, pres.slide_height
    slides = []
    for slide in pres.slides:
        texts, films = [], []
        for shape in slide.shapes:
            link = shape.element.xpath(".//a:videoFile/@r:link")
            if link:
                part = shape.part.rels[link[0]].target_part
                slug = by_size.get(part_size.get(str(part.partname)))
                assert slug, f"film on slide {len(slides) + 1} does not match a file in Projects/"
                films.append((slug, shape.left / w, shape.top / h, shape.width / w, shape.height / h))
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs).strip()
                    if (line and not FOOTER_NOISE.match(line) and not PRIVATE.search(line)
                            and "PLACEHOLDER" not in line and ".mp4" not in line):
                        texts.append(line)
        slides.append((texts, films))
    return slides


def slide_markup(n, total, texts, films):
    alt = f"Slide {n} of {total}. " + " · ".join(texts)
    if len(alt) > 420:
        alt = alt[:417].rsplit(" ", 1)[0] + "…"
    lazy = "" if n <= 2 else ' loading="lazy" decoding="async"'
    out = [f'<section class="slide" id="slide-{n}" aria-label="Slide {n} of {total}">',
           f'  <img src="assets/deck/slide-{n:02d}.webp" srcset="assets/deck/slide-{n:02d}.webp 1280w, '
           f'assets/deck/slide-{n:02d}-2x.webp 2560w" sizes="(min-width: 1408px) 1280px, 96vw" '
           f'width="1280" height="720"{lazy} alt="{html.escape(alt, quote=True)}">']
    for slug, x, y, fw, fh in films:
        small = " slide-film--sm" if fw < 0.4 else ""
        out.append(f'  <button class="slide-film{small}" type="button" data-film="assets/video/{slug}.mp4" '
                   f'style="left:{x:.4%};top:{y:.4%};width:{fw:.4%};height:{fh:.4%}" '
                   f'aria-label="Play the {FILMS[slug][0]} film">{PLAY}</button>')
    out.append("</section>")
    return "\n".join(out)


def build_deck():
    import fitz  # PyMuPDF

    slides = read_deck()
    total = len(slides)
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", PS_EXPORT], check=True,
                       env={**os.environ, "DARIA_PPTX": str(PPTX), "DARIA_OUT": str(tmp)})
        pngs = sorted((tmp / "png").glob("*.PNG"), key=lambda p: int(re.sub(r"\D", "", p.stem)))
        assert len(pngs) == total, f"PowerPoint exported {len(pngs)} slides, the deck has {total}"

        clear(DECK)
        for n, png in enumerate(pngs, 1):
            im = Image.open(png).convert("RGB")
            im.save(DECK / f"slide-{n:02d}-2x.webp", "WEBP", quality=84, method=6)
            im.resize((1280, 720), Image.LANCZOS).save(DECK / f"slide-{n:02d}.webp", "WEBP", quality=86, method=6)
        print("deck  ", total, "slides", f"{sum(f.stat().st_size for f in DECK.iterdir()) / 1e6:.1f} MB")

        doc = fitz.open(tmp / "deck.pdf")
        assert doc.page_count == total
        doc.set_metadata({"title": "DARIA · AI Creative Studio · Portfolio",
                          "author": "DARIA · AI Creative Studio",
                          "subject": "Campaign films and stills for brands, made with AI.",
                          "keywords": "AI creative studio, AI product film, AI campaign stills, lookbook"})
        PDF.parent.mkdir(parents=True, exist_ok=True)
        doc.save(PDF, garbage=4, deflate=True)
        doc.close()  # Windows cannot clear the temp folder while the file is open
        print("pdf   ", total, "pages", f"{PDF.stat().st_size / 1e6:.1f} MB")

    block = "\n\n".join(slide_markup(n, total, texts, films) for n, (texts, films) in enumerate(slides, 1))
    page = PORTFOLIO.read_text(encoding="utf-8")
    page, hits = re.subn(r"(<!-- slides:start[^>]*-->).*?(<!-- slides:end -->)",
                         lambda m: f"{m.group(1)}\n{block}\n{m.group(2)}", page, flags=re.S)
    assert hits == 1, "slides:start / slides:end markers not found in portfolio.html"
    size = f"{PDF.stat().st_size / 1e6:.0f} MB"
    page = re.sub(r'(data-deck="count">)[^<]*', rf"\g<1>{total}", page)
    page = re.sub(r'(data-deck="size">)[^<]*', rf"\g<1>{size}", page)
    PORTFOLIO.write_text(page, encoding="utf-8")
    print("page   portfolio.html updated ·", sum(len(f) for _, f in slides), "films placed")


# -------------------------------------------------------------------- og ---

def browser():
    for exe in (r"C:\Program Files\Google\Chrome\Application\chrome.exe",
                r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
                r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
                r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
                "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"):
        if Path(exe).exists():
            return exe
    for name in ("google-chrome", "chromium", "chrome", "msedge"):
        if shutil.which(name):
            return shutil.which(name)
    sys.exit("Chrome or Edge is needed to render the share image.")


OG_HTML = """<!DOCTYPE html><html><head><meta charset="utf-8">
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@500&family=Playfair+Display:wght@600&display=swap" rel="stylesheet">
<style>
  html, body { margin: 0; width: 1200px; height: 630px; overflow: hidden; }
  body { display: flex; background: #FAF6F0; color: #241C17; font-family: Inter, "Segoe UI", sans-serif; }
  .t { flex: 1; padding: 64px 56px 56px 72px; display: flex; flex-direction: column; justify-content: space-between; }
  .w { font: 600 132px/.92 "Playfair Display", Georgia, serif; letter-spacing: .01em; }
  .s { margin-top: 20px; font: 500 15px/1 Inter, sans-serif; letter-spacing: .42em; text-transform: uppercase; color: #B9552F; }
  .c { font: 600 38px/1.25 "Playfair Display", Georgia, serif; max-width: 16ch; }
  .r { width: 56px; height: 2px; background: #B9552F; margin: 0 0 26px; }
  .m { width: 420px; background: url("%s") center 30%% / cover; }
</style></head><body>
<div class="t"><div><div class="w">DARIA</div><div class="s">AI Creative Studio</div></div>
<div><div class="r"></div><div class="c">Your next campaign doesn’t need a shoot.</div></div></div>
<div class="m"></div></body></html>"""


def build_og():
    with tempfile.TemporaryDirectory() as tmp, tempfile.TemporaryDirectory() as profile:
        tmp = Path(tmp)
        hero = tmp / "hero.jpg"
        flat(Image.open(SRC / "Chilli Beans/Final/Product Ad (1).png")).save(hero, "JPEG", quality=90)
        (tmp / "og.html").write_text(OG_HTML % hero.name, encoding="utf-8")
        shot = tmp / "og.png"
        subprocess.run([browser(), "--headless=new", "--disable-gpu", "--hide-scrollbars",
                        f"--user-data-dir={profile}", "--virtual-time-budget=20000",
                        "--window-size=1200,630", f"--screenshot={shot}", (tmp / "og.html").as_uri()],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        Image.open(shot).convert("RGB").save(SITE / "assets" / "og.jpg", "JPEG", quality=88, optimize=True)
        print("og     site/assets/og.jpg")


# ----------------------------------------------------------------- icons ---

# The mark: a cream Playfair "D" on pine, over the short gold rule the deck
# uses under its wordmark. Gold only ever sits on pine, so the pairing holds.
ICON_HTML = """<!DOCTYPE html><html><head><meta charset="utf-8">
<link href="https://fonts.googleapis.com/css2?family=Playfair+Display:wght@700&display=block" rel="stylesheet">
<style>
  html, body { margin: 0; width: 512px; height: 512px; overflow: hidden; background: transparent; }
  .tile { position: relative; width: 512px; height: 512px; background: #1F4640; border-radius: %(radius)s; }
  .d { position: absolute; left: 0; right: 0; top: 50px; text-align: center;
       font: 700 360px/1 "Playfair Display", Georgia, serif; color: #FAF6F0; }
  .rule { position: absolute; left: 50%%; top: 404px; width: 132px; height: 16px; margin-left: -66px; background: #C99A3E; }
</style></head><body><div class="tile"><div class="d">D</div><div class="rule"></div></div></body></html>"""


def render_icon(radius):
    with tempfile.TemporaryDirectory() as tmp, tempfile.TemporaryDirectory() as profile:
        tmp = Path(tmp)
        (tmp / "icon.html").write_text(ICON_HTML % {"radius": radius}, encoding="utf-8")
        shot = tmp / "icon.png"
        subprocess.run([browser(), "--headless=new", "--disable-gpu", "--hide-scrollbars",
                        f"--user-data-dir={profile}", "--virtual-time-budget=20000",
                        "--default-background-color=00000000",
                        "--window-size=512,512", f"--screenshot={shot}", (tmp / "icon.html").as_uri()],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        return Image.open(shot).convert("RGBA")


def build_icons():
    out = SITE / "assets" / "icons"
    clear(out)
    tile = render_icon("22%")     # browser tabs, bookmarks, Android
    square = render_icon("0")     # iOS rounds the corners itself
    for size in (32, 192, 512):
        tile.resize((size, size), Image.LANCZOS).save(out / f"icon-{size}.png", optimize=True)
    square.convert("RGB").resize((180, 180), Image.LANCZOS).save(out / "apple-touch-icon.png", optimize=True)
    tile.save(SITE / "favicon.ico", sizes=[(16, 16), (32, 32), (48, 48)])
    (SITE / "site.webmanifest").write_text(
        '{\n'
        '  "name": "DARIA \u00b7 AI Creative Studio",\n'
        '  "short_name": "DARIA",\n'
        '  "icons": [\n'
        '    { "src": "/assets/icons/icon-192.png", "sizes": "192x192", "type": "image/png" },\n'
        '    { "src": "/assets/icons/icon-512.png", "sizes": "512x512", "type": "image/png" }\n'
        '  ],\n'
        '  "theme_color": "#1F4640",\n'
        '  "background_color": "#FAF6F0",\n'
        '  "display": "standalone"\n'
        '}\n', encoding="utf-8")
    print("icons  favicon.ico, assets/icons/, site.webmanifest")


STEPS = {"deck": build_deck, "images": build_images, "video": build_video, "og": build_og,
         "icons": build_icons}
EXTRA = {"frames": build_frames, "refs": build_refs}  # parts of the steps above, on their own

if __name__ == "__main__":
    wanted = sys.argv[1:] or list(STEPS)
    known = {**STEPS, **EXTRA}
    for name in wanted:
        if name not in known:
            sys.exit(f"unknown step {name!r} — choose from {', '.join(known)}")
    for name in wanted:
        known[name]()
